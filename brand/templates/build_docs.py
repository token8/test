"""Builds brand-docs-template.docx from ../STYLE_GUIDE.md and ../tokens.json.

Upload to Google Drive; it converts to a Google Doc with the brand paragraph styles.
In Docs, run Format > Paragraph styles > Options > Save as my default styles.
"""
import json, re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

HERE = Path(__file__).parent
tokens = json.loads((HERE.parent / "tokens.json").read_text())
H, B = tokens["font"]["heading"]["workspace"], tokens["font"]["body"]["workspace"]
C = {k: RGBColor.from_string(v.lstrip("#")) for k, v in tokens["color"].items() if isinstance(v, str) and v.startswith("#")}

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.2))

def style(name, font, size, bold=False, color="ink", before=0, after=6, line=1.3):
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.bold, st.font.color.rgb = font, Pt(size), bold, C[color]
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    rf = st.element.rPr.rFonts
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), font)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(attr), None)  # theme fonts would override the brand font
    pf = st.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    return st

style("Normal", B, 11)
style("Title", H, 26, bold=True, after=4, line=1.1)
style("Subtitle", B, 13, color="muted", after=18)
style("Heading 1", H, 18, bold=True, before=20, after=6, line=1.1)
style("Heading 2", H, 14, bold=True, before=14, after=4, line=1.1)
style("Heading 3", H, 11, bold=True, color="accent_text", before=10, after=2, line=1.1)
style("Caption", B, 9, color="muted")
style("Quote", B, 11, color="muted")
style("List Bullet", B, 11, after=3)
# Title style has a bottom border by default in python-docx's template; remove it.
ppr = doc.styles["Title"].element.get_or_add_pPr()
for b in ppr.findall(qn("w:pBdr")):
    ppr.remove(b)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")

def add_inline(par, text, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.color.rgb = C["accent_text"]
        elif part.startswith("["):
            r = par.add_run(re.match(r"\[([^\]]+)\]", part).group(1))
        elif part.startswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if bold:
            r.bold = True

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)

def add_table(rows):
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, val, bold=(i == 0))
            for r in p.runs:
                r.font.size = Pt(9.5)
            if i == 0:
                shade(cell, "F4F4F4")
    doc.add_paragraph()

doc.add_picture(str(HERE.parent / "logo" / "xlrte-color.png"), height=Cm(0.9))
doc.paragraphs[-1].paragraph_format.space_after = Pt(18)

lines = (HERE.parent / "STYLE_GUIDE.md").read_text().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("|"):
        block = []
        while i < len(lines) and lines[i].startswith("|"):
            block.append(lines[i]); i += 1
        add_table(block); continue
    if ln.startswith("# "):
        doc.add_paragraph(ln[2:], style="Title")
    elif ln.startswith("## "):
        doc.add_paragraph(ln[3:], style="Heading 1")
    elif ln.startswith("### "):
        doc.add_paragraph(ln[4:], style="Heading 2")
    elif ln.startswith("**") and ln.endswith("**"):
        doc.add_paragraph(ln.strip("*"), style="Heading 3")
    elif ln.startswith("> "):
        text = [ln[2:]]
        while i + 1 < len(lines) and lines[i + 1].startswith(">"):
            i += 1; text.append(lines[i].lstrip("> "))
        add_inline(doc.add_paragraph(style="Quote"), " ".join(t for t in text if t))
    elif re.match(r"^\d+\. ", ln) or ln.startswith("- "):
        numbered = bool(re.match(r"^\d+\. ", ln))
        text = re.sub(r"^(\d+\. |- )", "", ln)
        while i + 1 < len(lines) and lines[i + 1].startswith("  ") and lines[i + 1].strip():
            i += 1; text += " " + lines[i].strip()
        if numbered:  # literal numbers, so each list restarts at 1 in Docs too
            par = doc.add_paragraph()
            par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.6), Cm(-0.6)
            par.paragraph_format.space_after = Pt(3)
            add_inline(par, re.match(r"^\d+\. ", ln).group(0).replace(" ", "\t") + text)
        else:
            add_inline(doc.add_paragraph(style="List Bullet"), text)
    elif ln.strip():
        text = ln
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||>|- |\d+\. |\*\*)", lines[i + 1]):
            i += 1; text += " " + lines[i].strip()
        add_inline(doc.add_paragraph(), text)
    i += 1

out = HERE / "brand-docs-template.docx"
doc.save(out)
print("wrote", out)
